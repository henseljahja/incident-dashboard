#!/usr/bin/env python3
"""Build REF_USER and REF_CUSTOM_FIELD_ENUM seeds from the delivered workbook.

``redmine.users``, ``redmine.email_addresses`` and ``redmine.custom_field_enumerations``
were never delivered with the sample dumps, yet the report joins all three. Every enum
column and all three of Task 8's label rules depend on them, so they need a source rather
than a promise. The one artefact that *does* carry the labels is the report we are
reproducing: ``docs/reference/Detail Incident, Problem, SR 2025 - 2026.xlsx``, 4,627 rows
of already-resolved ``author`` / ``assigned_to`` / ``Developer_assigned`` and the six enum
columns.

OUTPUT IS GIT-IGNORED. It carries real customer and staff names. The tests use the
synthetic, tracked ``src/test/resources/fixtures/ref_*.csv`` instead.

What the workbook can and cannot give us
----------------------------------------
REF_USER: fully derivable. ``author_id`` pairs with ``author`` and ``assigned_to_id``
pairs with ``assigned_to`` on every row, so id -> label is a direct read.

The two label columns are NOT the same kind of value, and the source query is where that
comes from::

    IFNULL(email_addresses.address, users.lastname) AS assigned_to

So a label containing ``@`` is an email address and lands in ADDRESS; anything else is
the LASTNAME fallback and lands in LASTNAME. Writing both into one column would destroy
exactly the distinction Task 8's rules read.

``Developer_assigned`` has NO id column in the workbook -- only the rendered label -- so
it cannot key a REF_USER row. Those labels are reported at the end as a coverage gap
rather than guessed at.

REF_CUSTOM_FIELD_ENUM: the workbook has the LABELS but not the ``enum_value_id`` behind
them, and the id is what ``custom_values.value`` actually stores. Inventing surrogate ids
would produce a table that looks populated and joins to nothing.

Instead the ids are recovered by joining: ``07. redmine.journal_details.sql`` records the
numeric enum id each field was set to, ``06. redmine.journals.sql`` ties that to an issue,
and the workbook gives that same issue's rendered label. Latest journal wins, matching the
report's ``MAX(journal_id)`` semantics. Labels that never appear in the journal dump are
written to ``ref_custom_field_enum_unresolved.csv`` for a DBA to fill in from the real
``custom_field_enumerations`` table -- unresolved, never fabricated.

Usage
-----
    python3 scripts/extract-ref-tables.py [--out DIR]

Requires openpyxl.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
REFERENCE_DIR = os.path.join(REPO_ROOT, "docs", "reference")
WORKBOOK = os.path.join(REFERENCE_DIR, "Detail Incident, Problem, SR 2025 - 2026.xlsx")
JOURNALS_SQL = os.path.join(REFERENCE_DIR, "06. redmine.journals.sql")
JOURNAL_DETAILS_SQL = os.path.join(REFERENCE_DIR, "07. redmine.journal_details.sql")

# Workbook column -> custom_field_id, per CFG_CUSTOM_FIELD_MAP (V2__seed_config.sql).
ENUM_COLUMNS = {
    "TYPE_OF_INCIDENT": 41,
    "CATEGORY_INCIDENT": 77,
    "IMPACT_SEVERITY": 93,
    "URGENCY_SEVERITY": 94,
    "INSTITUSI": 101,
    "EXTERNAL_INTERNAL": 349,
}

# Values openpyxl / the exporter render for a blank cell.
BLANK = {None, "", "None", "NULL", "nan"}


def blank(value) -> bool:
    return value is None or str(value).strip() in BLANK


def read_workbook():
    try:
        import openpyxl
    except ImportError:
        sys.exit(
            "openpyxl is required: python3 -m pip install openpyxl\n"
            "(pandas.read_excel works too, but openpyxl is the smaller dependency.)"
        )
    if not os.path.exists(WORKBOOK):
        sys.exit(f"workbook not found: {WORKBOOK}")

    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    index = {name: i for i, name in enumerate(header)}
    return [dict(zip(header, row)) for row in rows], index


def extract_users(rows):
    """id -> {'address', 'lastname'}, merged across every row that mentions the id.

    A user can appear as an author on one row and an assignee on another, and the two
    columns are resolved by different rules upstream, so the merge keeps whichever field
    each occurrence actually supplied instead of letting the last row win.
    """
    users: dict[int, dict[str, str]] = {}
    unkeyed_developers: set[str] = set()

    for row in rows:
        for id_column, label_column in (("author_id", "author"),
                                        ("assigned_to_id", "assigned_to")):
            raw_id = row.get(id_column)
            if blank(raw_id):
                continue
            try:
                user_id = int(float(str(raw_id).strip()))
            except ValueError:
                continue
            entry = users.setdefault(user_id, {"address": "", "lastname": ""})
            label = row.get(label_column)
            if blank(label):
                continue
            label = str(label).strip()
            # The source query's IFNULL(email_addresses.address, users.lastname): an '@'
            # means the email row existed, anything else is the lastname fallback.
            if "@" in label:
                entry["address"] = entry["address"] or label
            else:
                entry["lastname"] = entry["lastname"] or label

        developer = row.get("Developer_assigned")
        if not blank(developer):
            unkeyed_developers.add(str(developer).strip())

    return users, unkeyed_developers


def derive_names(entry):
    """Split what we have into FIRSTNAME / LASTNAME / DISPLAY_NAME.

    All three columns exist because Task 8 reads three of them by three different rules.
    Where the workbook only carries an email, the local part is the best available name;
    it is derived, and marked as such by keeping DISPLAY_NAME identical to it rather than
    inventing a plausible-looking human name.
    """
    address = entry["address"]
    lastname = entry["lastname"]
    if lastname:
        parts = lastname.split()
        firstname = parts[0]
        display = lastname
    elif address:
        local = address.split("@", 1)[0]
        parts = re.split(r"[._-]+", local)
        firstname = parts[0]
        lastname = parts[-1] if len(parts) > 1 else parts[0]
        display = local
    else:
        return None
    return firstname, lastname, display


def parse_sql_tuples(path):
    """Yield the VALUES tuple of each INSERT in a HeidiSQL dump, fields already split.

    The dumps have an EMPTY table name (``INSERT INTO `` (...)``) -- see
    extract-fixtures.py, defect (a) -- so the caller supplies the column meaning.
    """
    if not os.path.exists(path):
        sys.exit(f"dump not found: {path}")
    with open(path, encoding="utf-8", errors="replace") as handle:
        text = handle.read()
    for statement in re.findall(r"VALUES\s*\((.*?)\);\s*$", text, re.M | re.S):
        yield split_sql_fields(statement)


def split_sql_fields(text):
    fields, current, in_quote, escaped = [], [], False, False
    for ch in text:
        if escaped:
            current.append(ch)
            escaped = False
        elif ch == "\\" and in_quote:
            escaped = True
        elif ch == "'":
            in_quote = not in_quote
        elif ch == "," and not in_quote:
            fields.append("".join(current).strip())
            current = []
        else:
            current.append(ch)
    fields.append("".join(current).strip())
    return fields


def recover_enum_ids(rows):
    """Recover enum_value_id -> label by joining the journal dumps to the workbook.

    journals: (id, journalized_id, journalized_type, user_id, notes, created_on, ...)
    journal_details: (id, journal_id, property, prop_key, old_value, value)
    """
    journal_to_issue = {}
    for fields in parse_sql_tuples(JOURNALS_SQL):
        if len(fields) < 3:
            continue
        try:
            journal_to_issue[int(fields[0])] = int(fields[1])
        except ValueError:
            continue

    # (issue_id, custom_field_id) -> (journal_id, enum_value_id); latest journal wins,
    # matching the report's MAX(journal_id) "last value wins" semantics.
    latest: dict[tuple[int, int], tuple[int, int]] = {}
    for fields in parse_sql_tuples(JOURNAL_DETAILS_SQL):
        if len(fields) < 6:
            continue
        try:
            journal_id = int(fields[1])
        except ValueError:
            continue
        prop = fields[2].strip()
        key = fields[3].strip()
        value = fields[5].strip()
        if prop != "cf" or not key.isdigit() or not value.isdigit():
            continue
        field_id = int(key)
        if field_id not in ENUM_COLUMNS.values():
            continue
        issue_id = journal_to_issue.get(journal_id)
        if issue_id is None:
            continue
        slot = (issue_id, field_id)
        if slot not in latest or latest[slot][0] < journal_id:
            latest[slot] = (journal_id, int(value))

    labels_by_issue = {}
    for row in rows:
        raw = row.get("No_incident")
        if blank(raw):
            continue
        try:
            labels_by_issue[int(float(str(raw).strip()))] = row
        except ValueError:
            continue

    resolved: dict[tuple[int, int], str] = {}
    all_labels: dict[int, set[str]] = {fid: set() for fid in ENUM_COLUMNS.values()}
    for column, field_id in ENUM_COLUMNS.items():
        for row in rows:
            label = row.get(column)
            if not blank(label):
                all_labels[field_id].add(str(label).strip())

    for (issue_id, field_id), (_, enum_id) in latest.items():
        row = labels_by_issue.get(issue_id)
        if row is None:
            continue
        column = next(c for c, f in ENUM_COLUMNS.items() if f == field_id)
        label = row.get(column)
        if blank(label):
            continue
        resolved[(field_id, enum_id)] = str(label).strip()

    unresolved = {
        field_id: sorted(labels - {v for (f, _), v in resolved.items() if f == field_id})
        for field_id, labels in all_labels.items()
    }
    return resolved, unresolved


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default=os.path.join(os.path.dirname(__file__), "..",
                                                      "target", "ref-seed"),
                        help="output directory (git-ignored; defaults to target/ref-seed)")
    args = parser.parse_args()
    out_dir = os.path.abspath(args.out)
    os.makedirs(out_dir, exist_ok=True)

    rows, _ = read_workbook()
    print(f"workbook: {len(rows)} data rows")

    users, unkeyed_developers = extract_users(rows)
    user_path = os.path.join(out_dir, "ref_user.csv")
    written = 0
    no_address = 0
    with open(user_path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["user_id", "address", "firstname", "lastname", "display_name"])
        for user_id in sorted(users):
            names = derive_names(users[user_id])
            if names is None:
                # An id that appears with no label anywhere. Emitting a row of NULLs would
                # be worse than omitting it: 'Unknown (<id>)' is the designed degradation.
                continue
            firstname, lastname, display = names
            address = users[user_id]["address"]
            if not address:
                no_address += 1
            writer.writerow([user_id, address or "\\N", firstname, lastname, display])
            written += 1
    print(f"REF_USER: {written} rows -> {user_path}"
          f"  ({no_address} with no email address, resolved by lastname)")
    print(f"  Developer_assigned labels with no id column in the workbook: "
          f"{len(unkeyed_developers)} (coverage gap; needs redmine.users to key)")

    resolved, unresolved = recover_enum_ids(rows)
    enum_path = os.path.join(out_dir, "ref_custom_field_enum.csv")
    with open(enum_path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["custom_field_id", "enum_value_id", "label", "position"])
        position = {}
        for (field_id, enum_id) in sorted(resolved):
            position[field_id] = position.get(field_id, 0) + 1
            writer.writerow([field_id, enum_id, resolved[(field_id, enum_id)],
                             position[field_id]])
    print(f"REF_CUSTOM_FIELD_ENUM: {len(resolved)} rows -> {enum_path}")

    gap_path = os.path.join(out_dir, "ref_custom_field_enum_unresolved.csv")
    total_gap = sum(len(v) for v in unresolved.values())
    with open(gap_path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["custom_field_id", "enum_value_id", "label", "position"])
        for field_id in sorted(unresolved):
            for label in unresolved[field_id]:
                # enum_value_id left EMPTY on purpose: the workbook does not carry it and
                # the journal dump never showed this label being set. A DBA fills it in
                # from redmine.custom_field_enumerations. Never guessed.
                writer.writerow([field_id, "", label, ""])
    print(f"UNRESOLVED labels (need custom_field_enumerations.id): {total_gap} -> {gap_path}")
    for field_id in sorted(unresolved):
        print(f"  custom_field_id {field_id}: "
              f"{len([1 for (f, _) in resolved if f == field_id])} resolved, "
              f"{len(unresolved[field_id])} unresolved")


if __name__ == "__main__":
    main()
